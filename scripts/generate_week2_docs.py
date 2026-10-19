"""Script to generate the official Week 2 submission Word document and submission text."""

from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets padding for a table cell in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def add_callout(doc, text, title="KEY DOMAIN INSIGHT"):
    """Adds a stylish callout box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    # Set border
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="1B365D"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f"[{title}] ")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    run_body = p.add_run(text)
    run_body.font.name = "Calibri"
    run_body.font.size = Pt(10)
    run_body.font.italic = True
    run_body.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def build_docx(output_path: Path, figures_dir: Path):
    doc = docx.Document()

    # Configure Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Palette definition
    NAVY = RGBColor(0x1B, 0x36, 0x5D)
    DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
    SLATE = RGBColor(0x55, 0x55, 0x55)

    # ==================== COVER / TITLE SECTION ====================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(6)
    title_run = title_p.add_run("Retail Sales Forecasting & Inventory Planning Using Python")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(24)
    title_run.bold = True
    title_run.font.color.rgb = NAVY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(24)
    sub_run = sub_p.add_run("Week 2 Deliverable: Exploratory Data Analysis, Data Quality Audit & Preparation")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = SLATE

    meta_tbl = doc.add_table(rows=4, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta_data = [
        ("Internship Program:", "YuvaIntern Virtual Data Science Explorer Internship"),
        ("Project Domain:", "Retail Time Series Forecasting & Supply Chain Inventory Optimization"),
        ("Technical Track:", "Artificial Intelligence & Data Science (B.Tech Practical Curriculum)"),
        ("Submission Status:", "Week 2 Verified Technical Report (Version 1.0)"),
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        c0, c1 = meta_tbl.cell(row_idx, 0), meta_tbl.cell(row_idx, 1)
        c0.text = k
        c0.paragraphs[0].runs[0].bold = True
        c0.paragraphs[0].runs[0].font.size = Pt(10)
        c0.paragraphs[0].runs[0].font.name = "Calibri"
        c0.paragraphs[0].runs[0].font.color.rgb = NAVY
        c1.text = v
        c1.paragraphs[0].runs[0].font.size = Pt(10)
        c1.paragraphs[0].runs[0].font.name = "Calibri"
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)

    doc.add_page_break()

    # ==================== SECTION 1: EXECUTIVE SUMMARY ====================
    h1 = doc.add_heading("1. Executive Summary", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph(
        "This technical deliverable documents the execution and empirical validation of Week 2 for the project "
        "'Retail Sales Forecasting and Inventory Planning Using Python', conducted under the YuvaIntern Virtual Data Science "
        "Explorer Internship. Following the successful submission of the Week 1 project charter and architecture roadmap, "
        "Week 2 establishes the empirical data foundation required for downstream predictive time-series modeling and supply "
        "chain replenishment planning."
    )
    p.paragraph_format.space_after = Pt(6)

    p2 = doc.add_paragraph(
        "The project utilizes historical retail sales records across 45 unique retail outlets observed over 143 consecutive weeks "
        "(spanning February 5, 2010 through October 26, 2012), providing a balanced panel of exactly 6,435 observations. A complete data "
        "integrity audit confirmed zero missing values and zero duplicate records. Rather than applying standard mathematical outlier truncation, "
        "a domain-specific outlier investigation established that all 34 statistical outlier weeks represent critical holiday demand surges "
        "(Thanksgiving/Black Friday and pre-Christmas shopping weeks). Preserving these peaks is paramount to preventing stockouts in "
        "downstream inventory planning."
    )
    p2.paragraph_format.space_after = Pt(6)

    p3 = doc.add_paragraph(
        "Exploratory Data Analysis confirmed a statistically significant holiday demand uplift of +7.84% (Welch's two-sample t-test: t = 2.6801, "
        "p = 0.00760; Mann-Whitney U test: U = 1,431,297, p = 0.02589), alongside extreme store-level heterogeneity where top-performing stores "
        "generate 8.11 times the revenue of bottom-tier stores. Modular scripts (download_data.py, data_cleaning.py, eda.py) and reproducible "
        "Jupyter notebooks have been deployed and validated in the workspace."
    )
    p3.paragraph_format.space_after = Pt(12)

    # ==================== SECTION 2: DATASET PROVENANCE & GOVERNANCE ====================
    doc.add_heading("2. Dataset Provenance, Scope & Governance", level=1)
    p = doc.add_paragraph(
        "Data science best practices require clear provenance verification and licensing compliance before undertaking predictive modeling. "
        "The characteristics and governance terms of the dataset are detailed below:"
    )
    p.paragraph_format.space_after = Pt(6)

    ds_tbl = doc.add_table(rows=6, cols=2)
    ds_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    ds_rows = [
        ("Benchmark Source", "Kaggle Competition: 'Walmart Recruiting - Store Sales Forecasting' (Walmart Inc., 2014)"),
        ("Intellectual Property", "Proprietary sales and operational data owned by Walmart Inc."),
        ("Licensing & Terms", "Released under Kaggle Competition Rules for educational benchmarking, academic study, and non-commercial research. Commercial redistribution is restricted."),
        ("Compliance Architecture", "In accordance with governance standards, raw dataset files are excluded from Git tracking via .gitignore. Fully reproducible automated retrieval is provided via src/download_data.py."),
        ("Public Mirror Source", "https://raw.githubusercontent.com/DrPoojaAbhijith/Walmart-Sales-Forecasting/main/Walmart%20DataSet.csv"),
        ("Dataset Dimensions", "6,435 rows × 8 original features (14 features post-engineering)"),
    ]
    for r_idx, (col1, col2) in enumerate(ds_rows):
        c0, c1 = ds_tbl.cell(r_idx, 0), ds_tbl.cell(r_idx, 1)
        c0.text = col1
        c0.paragraphs[0].runs[0].bold = True
        c0.paragraphs[0].runs[0].font.size = Pt(9.5)
        c0.paragraphs[0].runs[0].font.name = "Calibri"
        c1.text = col2
        c1.paragraphs[0].runs[0].font.size = Pt(9.5)
        c1.paragraphs[0].runs[0].font.name = "Calibri"
        set_cell_background(c0, "F8F9FA")
        set_cell_margins(c0, 70, 70, 100, 100)
        set_cell_margins(c1, 70, 70, 100, 100)

    # ==================== SECTION 3: DATA STRUCTURE & INTEGRITY AUDIT ====================
    doc.add_heading("3. Data Structure & Data Quality Audit", level=1)
    p = doc.add_paragraph(
        "A comprehensive structural audit was conducted by executing src/data_cleaning.py. The dataset contains eight attributes "
        "capturing sales, calendar flags, and regional environmental indicators. The audit confirmed 100% data completeness:"
    )
    p.paragraph_format.space_after = Pt(6)

    schema_tbl = doc.add_table(rows=9, cols=5)
    schema_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Column Name", "Raw Type", "Null Count", "Unique Values", "Operational Range / Meaning"]
    for c_idx, h in enumerate(headers):
        cell = schema_tbl.cell(0, c_idx)
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9.5)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, 80, 80, 80, 80)

    schema_data = [
        ("Store", "int64", "0 (0.0%)", "45", "Categorical identifier for retail outlet (1 to 45)."),
        ("Date", "object (str)", "0 (0.0%)", "143", "Weekly date step in DD-MM-YYYY format (2010-02-05 to 2012-10-26)."),
        ("Weekly_Sales", "float64", "0 (0.0%)", "6,435", "Weekly gross revenue in USD ($209,986.25 to $3,818,686.45). Target."),
        ("Holiday_Flag", "int64", "0 (0.0%)", "2", "Binary indicator (1 = designated holiday week, 0 = regular week)."),
        ("Temperature", "float64", "0 (0.0%)", "3,528", "Average regional ambient temperature in °F (-2.06°F to 100.14°F)."),
        ("Fuel_Price", "float64", "0 (0.0%)", "892", "Regional cost of fuel per gallon in USD ($2.472 to $4.468)."),
        ("CPI", "float64", "0 (0.0%)", "2,145", "Consumer Price Index inflation metric (126.06 to 227.23)."),
        ("Unemployment", "float64", "0 (0.0%)", "340", "Regional unemployment rate percentage (3.879% to 14.313%)."),
    ]
    for r_idx, row in enumerate(schema_data, start=1):
        for c_idx, val in enumerate(row):
            cell = schema_tbl.cell(r_idx, c_idx)
            cell.text = val
            cell.paragraphs[0].runs[0].font.size = Pt(9)
            cell.paragraphs[0].runs[0].font.name = "Calibri"
            if r_idx % 2 == 1:
                set_cell_background(cell, "F9FAFB")
            set_cell_margins(cell, 60, 60, 80, 80)

    p_audit = doc.add_paragraph()
    p_audit.paragraph_format.space_before = Pt(6)
    p_audit.add_run("Audit Highlights: ").bold = True
    p_audit.add_run(
        "Total duplicate rows: 0. Total missing entries: 0. Date granularity: exactly 143 contiguous weekly steps per store with no gaps. "
        "The panel is perfectly balanced across all 45 stores."
    )

    # ==================== SECTION 4: CLEANING & FEATURE ENGINEERING ====================
    doc.add_heading("4. Cleaning & Feature Engineering Methodology", level=1)
    p = doc.add_paragraph(
        "To prepare the raw data for chronological machine learning models and avoid data leakage, the following deterministic "
        "transformations were implemented in src/data_cleaning.py:"
    )
    p.paragraph_format.space_after = Pt(4)

    steps = [
        ("1. Explicit Date Parsing: ", "Raw date strings in DD-MM-YYYY format were explicitly parsed into datetime64[ns] to prevent month-day transposition errors."),
        ("2. Chronological Ordering: ", "Observations were sorted by Store and Date in ascending order, establishing a strict temporal progression necessary for walk-forward validation and lag engineering."),
        ("3. Calendar Feature Extraction: ", "Extracted Year, Month, WeekOfYear (ISO-8601), Day, and Quarter to allow models to learn recurring intra-year seasonal patterns."),
        ("4. Pre-Christmas Rush Indicator (Is_PreChristmas): ", "Identified and engineered a binary flag for Week 51 (December 17–24). In the raw dataset, Holiday_Flag is assigned to the final week of December (New Year's Eve), while actual retail transaction volume surges during the pre-Christmas shopping week."),
        ("5. Target Integrity Validation: ", "Verified that all Weekly_Sales values are strictly positive (> 0), confirming the absence of unhandled product return anomalies."),
    ]
    for title, desc in steps:
        p_step = doc.add_paragraph(style="List Bullet")
        p_step.paragraph_format.space_after = Pt(2)
        r1 = p_step.add_run(title)
        r1.bold = True
        r1.font.name = "Calibri"
        r2 = p_step.add_run(desc)
        r2.font.name = "Calibri"

    add_callout(
        doc,
        "By generating calendar indicators prior to modeling and establishing strict chronological sequence, "
        "we eliminate target leakage and ensure that validation splits mirror real-world operational forecasting conditions.",
        "FEATURE ENGINEERING PRINCIPLE"
    )

    # ==================== SECTION 5: OUTLIER ANALYSIS & DOMAIN POLICY ====================
    doc.add_heading("5. Outlier Diagnostics & Operational Justification", level=1)
    p = doc.add_paragraph(
        "Statistical outlier detection was performed using the standard Tukey Interquartile Range (IQR) method:"
    )
    p.paragraph_format.space_after = Pt(4)

    # IQR Table
    iqr_tbl = doc.add_table(rows=6, cols=2)
    iqr_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    iqr_metrics = [
        ("25th Percentile (Q1)", "$553,350.11"),
        ("75th Percentile (Q3)", "$1,420,158.66"),
        ("Interquartile Range (IQR)", "$866,808.55"),
        ("Lower Bound (Q1 - 1.5 × IQR)", "-$746,862.73 (Effective lower limit: $0.00)"),
        ("Upper Bound (Q3 + 1.5 × IQR)", "$2,720,371.49"),
        ("Identified Statistical Outliers", "34 records (0.5284% of total dataset)"),
    ]
    for r_idx, (k, v) in enumerate(iqr_metrics):
        c0, c1 = iqr_tbl.cell(r_idx, 0), iqr_tbl.cell(r_idx, 1)
        c0.text = k
        c0.paragraphs[0].runs[0].bold = True
        c0.paragraphs[0].runs[0].font.size = Pt(9.5)
        c1.text = v
        c1.paragraphs[0].runs[0].font.size = Pt(9.5)
        set_cell_background(c0, "F8F9FA")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)

    p_out = doc.add_paragraph()
    p_out.paragraph_format.space_before = Pt(6)
    p_out.add_run("Empirical Investigation of Outlier Timestamps:\n").bold = True
    p_out.add_run(
        "A granular breakdown of the 34 outlier records revealed that 100% of them occurred during peak Q4 promotional windows:\n"
        "• 2010-12-24 (Christmas Eve week): 9 stores (Peak: Store 14 at $3,818,686.45)\n"
        "• 2011-12-23 (Pre-Christmas week): 7 stores (Peak: Store 4 at $3,676,388.98)\n"
        "• 2010-11-26 (Thanksgiving / Black Friday): 5 stores\n"
        "• 2010-12-17 (Mid-December surge): 5 stores\n"
        "• 2011-11-25 (Thanksgiving / Black Friday): 4 stores\n"
        "• 2011-12-16 (Mid-December surge): 3 stores\n"
        "• 2010-12-10 (Early December shopping rush): 1 store\n"
    )

    add_callout(
        doc,
        "In generic data engineering, outliers are often winsorized or pruned. In retail operations, however, "
        "the Thanksgiving and Christmas surges generate a massive percentage of annual corporate profit. Pruning these records "
        "would bias predictive models toward under-forecasting peak demand, resulting in stockouts during the most lucrative "
        "weeks of the year. All 34 outliers were retained.",
        "SUPPLY CHAIN JUSTIFICATION"
    )

    # ==================== SECTION 6: EXPLORATORY DATA ANALYSIS ====================
    doc.add_heading("6. Exploratory Data Analysis & Empirical Visualizations", level=1)
    p = doc.add_paragraph(
        "Visual analytics were generated at 300 DPI using Matplotlib and Seaborn via src/eda.py. Below are four key visualizations "
        "illustrating seasonal dynamics, holiday impacts, store variance, and environmental correlations:"
    )
    p.paragraph_format.space_after = Pt(8)

    # Chart 1: Total Sales Trend
    doc.add_heading("Figure 1: Aggregate Nationwide Weekly Sales Trend (2010–2012)", level=2)
    fig1_path = figures_dir / "01_total_sales_trend.png"
    if fig1_path.exists():
        doc.add_picture(str(fig1_path), width=Inches(6.2))
        p_cap = doc.add_paragraph("Figure 1: Aggregate weekly retail revenue across all 45 stores with holiday annotations.")
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.italic = True
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_f1 = doc.add_paragraph(
        "Analysis of Figure 1: Nationwide sales demonstrate a remarkably consistent baseline oscillating between $45M and $50M per week. "
        "The dominant features of the series are the recurring annual November-December peaks, reaching nearly $80M per week—a surge of approximately "
        "60% over baseline. Secondary peaks occur during early February (Super Bowl) and early September (Labor Day)."
    )
    p_f1.paragraph_format.space_after = Pt(8)

    # Chart 2: Holiday Uplift
    doc.add_heading("Figure 2: Distributional Impact of Holiday Weeks on Store Sales", level=2)
    fig2_path = figures_dir / "02_holiday_vs_nonholiday_sales.png"
    if fig2_path.exists():
        doc.add_picture(str(fig2_path), width=Inches(6.2))
        p_cap = doc.add_paragraph("Figure 2: Boxplot and Kernel Density Estimate (KDE) comparing regular versus holiday weeks.")
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.italic = True
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_f2 = doc.add_paragraph(
        "Statistical Testing of Holiday Uplift: Mean weekly sales on regular weeks equal $1,041,256.38 (std: $558,957.44; n = 5,985), "
        "while mean sales on designated holiday weeks rise to $1,122,887.89 (std: $627,684.93; n = 450), reflecting an empirical uplift of +7.84%. "
        "To confirm that this uplift is not a sampling artifact, formal hypothesis tests were executed:\n"
        "• Welch's Two-Sample t-test (unequal variances): t = 2.6801, p-value = 0.00760 (significant at α = 0.01)\n"
        "• Mann-Whitney U test (non-parametric rank sum): U = 1,431,297.0, p-value = 0.02589 (significant at α = 0.05)\n"
        "Both tests confirm that holiday weeks produce statistically significant demand shifts, necessitating dedicated holiday indicators in predictive models."
    )
    p_f2.paragraph_format.space_after = Pt(8)

    # Chart 3: Store Tiers
    doc.add_heading("Figure 3: Cumulative Store Revenue Tiers & Performance Disparity", level=2)
    fig3_path = figures_dir / "03_store_sales_distribution.png"
    if fig3_path.exists():
        doc.add_picture(str(fig3_path), width=Inches(6.2))
        p_cap = doc.add_paragraph("Figure 3: Total cumulative revenue across all 45 outlets ranked by tier.")
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.italic = True
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_f3 = doc.add_paragraph(
        "Analysis of Figure 3: A profound structural divide exists across store locations. The top-performing outlet (Store 20) generated "
        "$301,397,792.46 in revenue over the period (mean: $2.11M/week), whereas the lowest-performing outlet (Store 33) generated $37,160,221.96 "
        "(mean: $260k/week). This represents an 8.11-fold volume disparity. Practical implication: applying uniform safety stock or global sales "
        "averages across stores will cause simultaneous severe over-stocking at small stores and chronic stockouts at flagship stores. Store-level "
        "segmentation is essential."
    )
    p_f3.paragraph_format.space_after = Pt(8)

    # Chart 4: Correlation Heatmap
    doc.add_heading("Figure 4: Feature Correlation Matrix (Sales & Macroeconomic Indicators)", level=2)
    fig4_path = figures_dir / "04_correlation_heatmap.png"
    if fig4_path.exists():
        doc.add_picture(str(fig4_path), width=Inches(5.5))
        p_cap = doc.add_paragraph("Figure 4: Pearson correlation matrix between sales, calendar features, and macroeconomic variables.")
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.italic = True
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_f4 = doc.add_paragraph(
        "Analysis of Figure 4: Direct linear correlations between Weekly_Sales and environmental regressors are weak (|r| < 0.15 for CPI, "
        "Unemployment, Temperature, and Fuel Price). In contrast, temporal features (WeekOfYear, Month, Is_PreChristmas) exhibit strong "
        "predictive alignment with demand spikes. This finding demonstrates that linear models relying solely on macroeconomic variables will fail; "
        "effective forecasting requires calendar feature engineering and non-linear tree-based architectures."
    )
    p_f4.paragraph_format.space_after = Pt(12)

    # ==================== SECTION 7: PRACTICAL INVENTORY SIGNIFICANCE ====================
    doc.add_heading("7. Practical Significance for Inventory Planning", level=1)
    p = doc.add_paragraph(
        "The empirical findings from Week 2 directly determine the parameterization of the supply chain inventory models to be developed in Week 5:"
    )
    p.paragraph_format.space_after = Pt(4)

    inv_points = [
        ("Safety Stock Buffer Formulation: ", "Safety stock is calculated as SS = Z × σ_L, where Z is the service level factor (e.g., Z = 1.65 for 95% cycle service level) and σ_L is lead-time demand standard deviation. Because demand variance expands dramatically during holiday periods, safety stock buffers must dynamically scale in November and December to prevent stockouts."),
        ("Reorder Point (ROP) Dynamics: ", "The reorder threshold ROP = d_L + SS (where d_L is expected demand over lead time) must transition from a static formula to a time-varying schedule that accounts for store-specific sales tiers and holiday demand multipliers."),
        ("Multi-Tier Inventory Allocation: ", "Given the 8.11× throughput differential between Store 20 and Store 33, distribution center replenishment schedules must prioritize Tier-1 high-velocity stores with more frequent delivery cycles, while low-velocity outlets operate on longer replenishment cycles to minimize holding costs."),
    ]
    for k, v in inv_points:
        p_inv = doc.add_paragraph(style="List Bullet")
        p_inv.paragraph_format.space_after = Pt(3)
        r1 = p_inv.add_run(k)
        r1.bold = True
        r2 = p_inv.add_run(v)

    # ==================== SECTION 8: LIMITATIONS ====================
    doc.add_heading("8. Limitations & Methodological Constraints", level=1)
    lims = [
        ("Aggregation Level: ", "Data is aggregated at the store-wide weekly revenue level. Department-level and SKU-level (item barcode) granularity is unavailable in this table, meaning inventory buffers apply at aggregate category replenishment levels."),
        ("Macroeconomic Frequency: ", "CPI and Unemployment series change on a monthly or quarterly basis rather than weekly, limiting their ability to explain week-to-week sales fluctuations."),
        ("Historical Temporal Window: ", "The 2.75-year timeline (2010–2012) provides two complete seasonal cycles. While sufficient for 4-to-12-week forward projections, long-term multi-decade structural shifts cannot be evaluated."),
    ]
    for k, v in lims:
        p_l = doc.add_paragraph(style="List Bullet")
        p_l.paragraph_format.space_after = Pt(2)
        p_l.add_run(k).bold = True
        p_l.add_run(v)

    # ==================== SECTION 9: CONCLUSION & ROADMAP ====================
    doc.add_heading("9. Conclusion & Weeks 3–6 Roadmap", level=1)
    p_c = doc.add_paragraph(
        "Week 2 data preparation and exploratory analysis has established a clean, validated foundation for the remaining internship tasks. "
        "The project roadmap proceeds as follows:"
    )
    p_c.paragraph_format.space_after = Pt(4)

    roadmap_steps = [
        ("Week 3: Baseline Forecasting: ", "Establish naive lag, moving average, and seasonal baseline forecasts using chronological train/test splits (2010–2011 train, 2012 test) to prevent target leakage."),
        ("Week 4: Advanced Predictive Modeling: ", "Train Ridge/Lasso Regressors and ensemble architectures (Random Forest, Gradient Boosting) incorporating calendar features, store embeddings, and holiday flags."),
        ("Week 5: Model Evaluation & Inventory Formulas: ", "Evaluate forecasts via MAE, RMSE, and WMAPE. Translate forecast errors into dynamic Safety Stock, Lead Time Demand, and Reorder Point schedules."),
        ("Week 6: Final Reporting & Recommendations: ", "Synthesize findings into an executive presentation and evidence-based supply chain operational strategy."),
    ]
    for k, v in roadmap_steps:
        p_r = doc.add_paragraph(style="List Bullet")
        p_r.paragraph_format.space_after = Pt(2)
        p_r.add_run(k).bold = True
        p_r.add_run(v)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"[SUCCESS] Word document generated at: {output_path}")


def build_submission_description(output_path: Path):
    text = (
        "YuvaIntern Virtual Data Science Explorer Internship - Week 2 Project Submission\n"
        "Project: Retail Sales Forecasting and Inventory Planning Using Python\n"
        "Repository: https://github.com/arikrishna-03/retail-sales-forecasting\n"
        "\n"
        "Submission Description:\n"
        "--------------------------------------------------------------------------------\n"
        "Following the successful submission and approval of the Week 1 project charter, "
        "Week 2 execution focused on the end-to-end data ingestion, data quality audit, cleaning, "
        "feature engineering, and exploratory data analysis (EDA) for historical weekly retail sales "
        "across 45 stores spanning February 2010 through October 2012 (6,435 balanced observations).\n"
        "\n"
        "1. Data Integrity and Audit:\n"
        "An initial structural audit validated 100% data completeness with exactly 0 missing values "
        "across all columns and 0 duplicate records. All 45 store outlets contain an identical panel of "
        "143 contiguous weekly observations without temporal gaps. In accordance with open-source data "
        "governance and Kaggle competition licensing rules, raw data files are safely excluded from Git "
        "version control via .gitignore, while an automated downloader script (src/download_data.py) "
        "guarantees complete end-to-end reproducibility.\n"
        "\n"
        "2. Cleaning and Feature Engineering:\n"
        "In src/data_cleaning.py, dates were explicitly parsed to datetime64[ns] from DD-MM-YYYY strings and "
        "chronologically sequenced by store to eliminate target leakage. Calendar features were extracted "
        "(Year, Month, WeekOfYear, Day, Quarter), alongside a specialized indicator (Is_PreChristmas) "
        "capturing Week 51 footfall surges which precede official calendar holiday markers.\n"
        "\n"
        "3. Outlier Diagnostics and Domain Decisions:\n"
        "Application of Tukey's 1.5*IQR threshold flagged 34 observations (0.53% of data) as mathematical "
        "outliers. An empirical timestamp audit proved that 100% of these records coincide with Thanksgiving "
        "(Black Friday) and Christmas shopping weeks (e.g., Store 14 peak sales of $3.82M on Dec 24, 2010). "
        "In retail supply chains, deleting peak demand spikes causes severe model under-prediction and costly "
        "stockouts; therefore, all outlier records were preserved and modeled via holiday features.\n"
        "\n"
        "4. Exploratory Data Analysis and Key Findings:\n"
        "Through src/eda.py, five publication-grade visualizations (300 DPI) were generated in reports/figures/:\n"
        "- Holiday Demand Uplift: Non-holiday weekly sales average $1,041,256.38, whereas holiday sales average "
        "$1,122,887.89 (+7.84% uplift). Statistical significance was formally proven via Welch's two-sample "
        "t-test (t = 2.6801, p = 0.00760) and the Mann-Whitney U test (U = 1,431,297, p = 0.02589).\n"
        "- Store Heterogeneity: Severe operational disparity exists between top-tier outlets (Store 20: $301.40M "
        "total revenue) and bottom-tier outlets (Store 33: $37.16M total revenue), an 8.11-fold difference.\n"
        "- Environmental Correlations: External variables (CPI, Unemployment, Temperature, Fuel Price) exhibit "
        "weak linear correlation with sales (|r| < 0.15), proving that calendar seasonality and store identity "
        "are the dominant drivers of demand.\n"
        "\n"
        "5. Inventory Planning Impact:\n"
        "These empirical findings directly guide Weeks 3-5 modeling: safety stock buffers (SS = Z * sigma_L) must "
        "dynamically scale during holiday spikes, while reorder points (ROP = d_L + SS) must be parameterized "
        "by store revenue tier rather than a naive global average.\n"
        "\n"
        "All code, notebooks (notebooks/01_data_inspection.ipynb, notebooks/02_exploratory_data_analysis.ipynb), "
        "and documentation are validated, reproducible, and ready for evaluation."
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(text, encoding="utf-8")
    words = len(text.split())
    print(f"[SUCCESS] Submission description generated at: {output_path} ({words} words)")


if __name__ == "__main__":
    project_root = Path(__file__).resolve().parent.parent
    reports_dir = project_root / "reports"
    figures_dir = reports_dir / "figures"

    docx_path = reports_dir / "Week_2_EDA_and_Data_Preparation.docx"
    build_docx(docx_path, figures_dir)

    txt_path = reports_dir / "Week_2_Submission_Description.txt"
    build_submission_description(txt_path)
